"""自定义表单插件：后台管理。

路由挂在核心 admin_bp 上（endpoint 归入 admin.*，自动获得后台地址前缀
即时生效机制），路径与核心版一致（/forms、/forms/<id>/edit、
/forms/<id>/submissions 等），老站书签/审计记录不受影响；
未启用插件时所有路由 404（不暴露存在性），菜单本就隐藏。

审计模块代码沿用核心版（'form' / 'form_submission'），历史审计日志在
插件启用后自动正常翻译显示。
"""
import io
import json
from datetime import datetime
from functools import wraps

from flask import (
    render_template, redirect, url_for, request,
    flash, abort, send_file
)
from openpyxl import Workbook

from app.extensions import db
from app.admin import admin_bp
from app.models.audit import (
    OP_CREATE, OP_UPDATE, OP_DELETE, OP_EXPORT, OP_BATCH,
)
from app.utils.helpers import permission_required, audit_log
from app.utils.uploads import save_upload_file
from app.plugin_system import plugin_enabled

from .models import Form, FormField, FormSubmission, FormSubmissionValue

# 审计模块代码（与核心版 app.models.audit 常量值一致，历史日志无缝翻译）
MODULE_FORM = 'form'
MODULE_FORM_SUBMISSION = 'form_submission'

FIELD_TYPES = [
    ('text', '单行文本'),
    ('textarea', '多行文本'),
    ('phone', '手机号'),
    ('email', '邮箱'),
    ('select', '下拉选择'),
    ('checkbox', '多选'),
    ('radio', '单选'),
    ('file', '文件上传'),
]


def _gate(view):
    """插件启用守卫：未启用 → 404。"""
    @wraps(view)
    def wrapper(*args, **kwargs):
        if not plugin_enabled('form'):
            abort(404)
        return view(*args, **kwargs)
    return wrapper


# ============ 表单管理 ============

@admin_bp.route('/forms')
@_gate
@permission_required('form:view')
def form_index():
    forms = Form.query.filter_by(is_deleted=False).order_by(Form.created_at.desc()).all()
    return render_template('admin/form/index.html', forms=forms)


@admin_bp.route('/forms/create', methods=['GET', 'POST'])
@_gate
@permission_required('form:manage')
def form_create():
    if request.method == 'POST':
        form = _save_form(None)
        if form is None:
            return redirect(url_for('admin.form_create'))
        return redirect(url_for('admin.form_edit', fid=form.id))
    return render_template('admin/form/form.html', form=None, field_types=FIELD_TYPES)


@admin_bp.route('/forms/<int:fid>/edit', methods=['GET', 'POST'])
@_gate
@permission_required('form:manage')
def form_edit(fid):
    form = Form.query.get_or_404(fid)
    if form.is_deleted:
        abort(404)
    if request.method == 'POST':
        updated = _save_form(form)
        if updated is None:
            return redirect(url_for('admin.form_edit', fid=fid))
        return redirect(url_for('admin.form_edit', fid=fid))
    fields = form.fields.filter_by(is_deleted=False).order_by(FormField.sort_order.asc()).all()
    return render_template('admin/form/form.html', form=form, fields=fields, field_types=FIELD_TYPES)


def _save_form(form):
    name = (request.form.get('name') or '').strip()
    slug = (request.form.get('slug') or '').strip().lower()
    if not name or not slug:
        flash('名称和标识必填', 'danger')
        return None

    existing = Form.query.filter_by(slug=slug, is_deleted=False).first()
    if existing and (form is None or existing.id != form.id):
        flash('表单标识已存在', 'danger')
        return None

    is_new = form is None
    if is_new:
        form = Form(slug=slug)
        db.session.add(form)

    form.name = name
    form.description = (request.form.get('description') or '').strip()
    form.success_message = (request.form.get('success_message') or '提交成功，感谢您的反馈！').strip()
    form.is_open = (request.form.get('is_open') == 'on')
    try:
        form.submit_interval = max(int(request.form.get('submit_interval') or 0), 0)
    except ValueError:
        form.submit_interval = 60

    db.session.flush()

    # 保存字段
    _save_form_fields(form)

    db.session.commit()
    flash('表单已保存', 'success')
    audit_log(OP_CREATE if is_new else OP_UPDATE, MODULE_FORM, form.id, form.name,
              {'slug': form.slug})
    return form


def _save_form_fields(form):
    keys = request.form.getlist('field_key[]')
    labels = request.form.getlist('field_label[]')
    types = request.form.getlist('field_type[]')
    requireds = request.form.getlist('field_required[]')
    placeholders = request.form.getlist('field_placeholder[]')
    helps = request.form.getlist('field_help[]')
    options_list = request.form.getlist('field_options[]')
    exts_list = request.form.getlist('field_exts[]')
    maxsizes = request.form.getlist('field_maxsize[]')
    ids = request.form.getlist('field_id[]')
    sorts = request.form.getlist('field_sort[]')

    keep_ids = set()
    seen_keys = set()
    keep_fields = []
    for i, key in enumerate(keys):
        key = (key or '').strip()
        label = (labels[i] if i < len(labels) else '').strip()
        ftype = types[i] if i < len(types) else 'text'
        if not key or not label:
            continue
        if key in seen_keys:
            flash(f'字段标识 {key} 重复，已忽略', 'warning')
            continue
        seen_keys.add(key)

        fid = int(ids[i]) if i < len(ids) and ids[i].isdigit() else None
        if fid:
            field = FormField.query.get(fid)
            if field is None or field.form_id != form.id:
                continue
            keep_ids.add(fid)
        else:
            field = FormField(form_id=form.id, field_key=key)
            db.session.add(field)

        field.label = label
        field.field_key = key
        field.field_type = ftype
        field.is_required = str(i) in requireds
        field.placeholder = (placeholders[i] if i < len(placeholders) else '').strip()
        field.help_text = (helps[i] if i < len(helps) else '').strip()
        field.sort_order = int(sorts[i]) if i < len(sorts) and sorts[i].isdigit() else i

        # 选项
        if ftype in ('select', 'checkbox', 'radio'):
            options_raw = (options_list[i] if i < len(options_list) else '').strip()
            options = [o.strip() for o in options_raw.split('\n') if o.strip()]
            field.options = json.dumps(options, ensure_ascii=False)
        else:
            field.options = None

        # 文件上传
        if ftype == 'file':
            field.allowed_exts = (exts_list[i] if i < len(exts_list) else '').strip()
            try:
                field.max_size = int(maxsizes[i]) * 1024 if i < len(maxsizes) and maxsizes[i].isdigit() else None
            except ValueError:
                field.max_size = None
        else:
            field.allowed_exts = None
            field.max_size = None

        field.is_deleted = False
        keep_fields.append(field)

    # 确保 ID 已生成
    db.session.flush()
    for f in keep_fields:
        if f.id:
            keep_ids.add(f.id)

    # 软删除未保留的字段（只处理本次未涉及的旧字段）
    for f in form.fields.all():
        if f.id and f.id not in keep_ids and not f.is_deleted:
            f.is_deleted = True


@admin_bp.route('/forms/<int:fid>/delete', methods=['POST'])
@_gate
@permission_required('form:manage')
def form_delete(fid):
    form = Form.query.get_or_404(fid)
    form.is_deleted = True
    db.session.commit()
    flash('表单已删除', 'success')
    audit_log(OP_DELETE, MODULE_FORM, form.id, form.name, {})
    return redirect(url_for('admin.form_index'))


# ============ 表单数据管理 ============

@admin_bp.route('/forms/<int:fid>/submissions')
@_gate
@permission_required('form:view')
def form_submissions(fid):
    form = Form.query.get_or_404(fid)
    page = max(int(request.args.get('page', 1)), 1)
    is_read = request.args.get('is_read', '')

    query = FormSubmission.query.filter_by(form_id=fid, is_deleted=False)
    if is_read == 'unread':
        query = query.filter_by(is_read=False)
    elif is_read == 'read':
        query = query.filter_by(is_read=True)

    pagination = query.order_by(FormSubmission.created_at.desc()).paginate(
        page=page, per_page=15, error_out=False
    )
    fields = form.fields.filter_by(is_deleted=False).order_by(FormField.sort_order.asc()).all()
    return render_template(
        'admin/form/submissions.html',
        form=form, submissions=pagination.items, pagination=pagination,
        fields=fields, is_read=is_read
    )


@admin_bp.route('/forms/<int:fid>/submissions/<int:sid>')
@_gate
@permission_required('form:view')
def form_submission_detail(fid, sid):
    sub = FormSubmission.query.get_or_404(sid)
    if sub.form_id != fid or sub.is_deleted:
        abort(404)
    if not sub.is_read:
        sub.is_read = True
        db.session.commit()
    fields = sub.form.fields.filter_by(is_deleted=False).order_by(FormField.sort_order.asc()).all()
    return render_template('admin/form/submission_detail.html', sub=sub, fields=fields)


@admin_bp.route('/forms/<int:fid>/submissions/<int:sid>/toggle-read', methods=['POST'])
@_gate
@permission_required('form:view')
def form_submission_toggle_read(fid, sid):
    sub = FormSubmission.query.get_or_404(sid)
    if sub.form_id != fid:
        abort(404)
    sub.is_read = not sub.is_read
    db.session.commit()
    return redirect(url_for('admin.form_submissions', fid=fid))


@admin_bp.route('/forms/<int:fid>/submissions/<int:sid>/delete', methods=['POST'])
@_gate
@permission_required('form:manage')
def form_submission_delete(fid, sid):
    sub = FormSubmission.query.get_or_404(sid)
    sub.is_deleted = True
    db.session.commit()
    flash('已删除', 'success')
    return redirect(url_for('admin.form_submissions', fid=fid))


@admin_bp.route('/forms/<int:fid>/submissions/batch', methods=['POST'])
@_gate
@permission_required('form:manage')
def form_submission_batch(fid):
    action = request.form.get('action')
    ids = [int(i) for i in request.form.getlist('ids[]') if i.isdigit()]
    if not ids:
        flash('未选择记录', 'warning')
        return redirect(url_for('admin.form_submissions', fid=fid))

    subs = FormSubmission.query.filter(FormSubmission.id.in_(ids), FormSubmission.form_id == fid).all()
    if action == 'read':
        for s in subs: s.is_read = True
    elif action == 'unread':
        for s in subs: s.is_read = False
    elif action == 'delete':
        for s in subs: s.is_deleted = True
    db.session.commit()
    flash('批量操作完成', 'success')
    return redirect(url_for('admin.form_submissions', fid=fid))


@admin_bp.route('/forms/<int:fid>/export')
@_gate
@permission_required('form:view')
def form_export(fid):
    """导出表单数据为 Excel。"""
    form = Form.query.get_or_404(fid)
    fields = form.fields.filter_by(is_deleted=False).order_by(FormField.sort_order.asc()).all()
    subs = FormSubmission.query.filter_by(form_id=fid, is_deleted=False).order_by(
        FormSubmission.created_at.desc()
    ).all()

    wb = Workbook()
    ws = wb.active
    ws.title = '表单数据'

    headers = ['ID', '提交时间', 'IP', '已读']
    for f in fields:
        headers.append(f.label)
    ws.append(headers)

    for sub in subs:
        row = [sub.id, sub.created_at.strftime('%Y-%m-%d %H:%M:%S'), sub.ip, '是' if sub.is_read else '否']
        for f in fields:
            val = sub.get_value(f.id)
            if f.field_type == 'file' and val:
                # 文件字段记录 URL
                row.append(val)
            else:
                row.append(val or '')
        ws.append(row)

    # 设置列宽
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else 'A' + chr(64 + i - 26)].width = 20

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f'{form.slug}_submissions_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    audit_log(OP_EXPORT, MODULE_FORM_SUBMISSION, None, form.name,
              {'form_id': fid, 'rows': len(subs), 'filename': filename})
    return send_file(
        buf,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
